import os
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from helpers import calc_pending, calc_received, calc_status

HEADERS = [
    'Date', 'Vehicle Number', 'State', 'City',
    'Driver Phone', 'Owner Phone',
    'Loading Address', 'Unloading Address',
    'Total Booking (₹)', 'Payment 1 (₹)', 'Payment 2 (₹)', 'Payment 3 (₹)',
    'Material', 'Weight (T)',
    'Received (₹)', 'Pending (₹)', 'Status',
    'Party Rate (₹)', 'Commission (₹)',
]

FIELD_MAP = [
    'date', 'vehicle_number', 'state', 'city',
    'driver_phone', 'owner_phone',
    'loading_address', 'unloading_address',
    'total_booking', 'payment_1', 'payment_2', 'payment_3',
]


def export_to_excel(trips, dest_dir=None):
    if dest_dir is None:
        dest_dir = os.path.expanduser('~/Desktop')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'TripLog'

    header_fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for row_idx, trip in enumerate(trips, start=2):
        # Columns 1–12: direct fields
        for col_idx, field in enumerate(FIELD_MAP, start=1):
            ws.cell(row=row_idx, column=col_idx, value=trip.get(field, ''))

        # Column 13: Material
        ws.cell(row=row_idx, column=13, value=trip.get('material', ''))

        # Column 14: Weight (T) — convert kg → tonnes
        weight_kg = trip.get('material_weight', 0) or 0
        ws.cell(row=row_idx, column=14, value=round(weight_kg / 1000, 3))

        # Columns 15–17: computed
        p1 = trip.get('payment_1', 0) or 0
        p2 = trip.get('payment_2', 0) or 0
        p3 = trip.get('payment_3', 0) or 0
        total = trip.get('total_booking', 0) or 0
        received = calc_received(p1, p2, p3)
        pending = calc_pending(total, p1, p2, p3)
        status = calc_status(pending, p1, p2, p3)

        ws.cell(row=row_idx, column=15, value=received)
        ws.cell(row=row_idx, column=16, value=pending)
        ws.cell(row=row_idx, column=17, value=status)

        # Columns 18–19: party rate and commission
        party_rate = trip.get('party_rate', 0) or 0
        commission = party_rate - total if party_rate > 0 else 0
        ws.cell(row=row_idx, column=18, value=party_rate)
        ws.cell(row=row_idx, column=19, value=commission)

    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    filename = f'TripLog_Export_{date.today().isoformat()}.xlsx'
    out_path = os.path.join(dest_dir, filename)
    wb.save(out_path)
    return out_path