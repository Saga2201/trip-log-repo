import os
from pathlib import Path
import openpyxl
from export import export_to_excel


SAMPLE_TRIPS = [
    {
        'id': 1, 'date': '2026-05-15', 'vehicle_number': 'GJ 12 AB 1234',
        'state': 'Gujarat', 'city': 'Surat',
        'driver_phone': '9876543210', 'owner_phone': '9123456789',
        'loading_address': 'Surat Market', 'unloading_address': 'Mumbai Hub',
        'total_booking': 18000.0, 'payment_1': 10000.0,
        'payment_2': 5000.0, 'payment_3': 3000.0,
        'created_at': '2026-05-15T10:00:00',
    },
    {
        'id': 2, 'date': '2026-05-10', 'vehicle_number': 'MH 04 XY 5678',
        'state': 'Maharashtra', 'city': 'Pune',
        'driver_phone': '', 'owner_phone': '',
        'loading_address': 'Pune Station', 'unloading_address': 'Delhi Hub',
        'total_booking': 32000.0, 'payment_1': 20000.0,
        'payment_2': 0.0, 'payment_3': 0.0,
        'created_at': '2026-05-10T09:00:00',
    },
]


def test_export_creates_xlsx_file(tmp_path):
    out_path = export_to_excel(SAMPLE_TRIPS, dest_dir=str(tmp_path))
    assert out_path.endswith('.xlsx')
    assert os.path.exists(out_path)


def test_export_correct_row_count(tmp_path):
    out_path = export_to_excel(SAMPLE_TRIPS, dest_dir=str(tmp_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    # header row + 2 data rows
    assert ws.max_row == 3


def test_export_includes_pending_and_status_columns(tmp_path):
    out_path = export_to_excel(SAMPLE_TRIPS, dest_dir=str(tmp_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    assert 'Pending (₹)' in headers
    assert 'Status' in headers


def test_export_pending_values_correct(tmp_path):
    out_path = export_to_excel(SAMPLE_TRIPS, dest_dir=str(tmp_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    pending_col = headers.index('Pending (₹)') + 1
    status_col = headers.index('Status') + 1
    # row 2 = first trip: 18000 - 10000 - 5000 - 3000 = 0, Paid
    assert ws.cell(2, pending_col).value == 0.0
    assert ws.cell(2, status_col).value == 'Paid'
    # row 3 = second trip: 32000 - 20000 = 12000, Partial
    assert ws.cell(3, pending_col).value == 12000.0
    assert ws.cell(3, status_col).value == 'Partial'


def test_export_includes_material_and_weight_columns(tmp_path):
    trips_with_material = [
        {**SAMPLE_TRIPS[0], 'material': 'Cotton Bales', 'material_weight': 8500.0},
        {**SAMPLE_TRIPS[1], 'material': '',             'material_weight': 0.0},
    ]
    out_path = export_to_excel(trips_with_material, dest_dir=str(tmp_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    assert 'Material' in headers
    assert 'Weight (T)' in headers


def test_export_weight_in_tonnes(tmp_path):
    trips_with_material = [
        {**SAMPLE_TRIPS[0], 'material': 'Steel Rods', 'material_weight': 12.5},
        {**SAMPLE_TRIPS[1], 'material': '',            'material_weight': 0.0},
    ]
    out_path = export_to_excel(trips_with_material, dest_dir=str(tmp_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    weight_col = headers.index('Weight (T)') + 1
    material_col = headers.index('Material') + 1
    # weight is stored in tonnes — exported as-is
    assert ws.cell(2, weight_col).value == 12.5
    assert ws.cell(2, material_col).value == 'Steel Rods'
    # row 3 = second trip: 0 T
    assert ws.cell(3, weight_col).value == 0.0
